"""表格渲染：生成**真正的 docx 表格对象**，绝不用制表符/空格堆文本。

必须做到的事（都是"看起来是表格"和"真的是表格"的区别）：
1. 用 `doc.add_table` 产出 `w:tbl`；
2. 边框宽度/颜色、单元格内边距、表头加粗与底色全部来自 `StyleSpec`；
3. 表头行可跨页重复（长表格断页后仍带表头）；
4. 本项目表格无合并单元格 —— 即便上游误报，这里也一律按矩形表渲染。
"""

from __future__ import annotations

from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from qqe.domain.models import Table
from qqe.domain.style import Align, StyleSpec, TableStyle
from qqe.renderer.fonts import apply_role_font
from qqe.renderer.paragraphs import add_paragraph, format_paragraph
from qqe.renderer.roles import Role, resolve

_TABLE_ALIGN = {
    Align.LEFT: WD_TABLE_ALIGNMENT.LEFT,
    Align.CENTER: WD_TABLE_ALIGNMENT.CENTER,
    Align.RIGHT: WD_TABLE_ALIGNMENT.RIGHT,
    Align.JUSTIFY: WD_TABLE_ALIGNMENT.CENTER,  # 表格没有"两端对齐"，退化为居中
}


def _twips(points: float) -> str:
    return str(round(points * 20))


def set_table_borders(table, *, width_pt: float, color: str) -> None:
    """边框粗细单位是 1/8 磅；颜色不带 #。"""
    tbl_pr = table._tbl.tblPr
    existing = tbl_pr.find(qn("w:tblBorders"))
    if existing is not None:
        tbl_pr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single" if width_pt > 0 else "none")
        element.set(qn("w:sz"), str(max(0, round(width_pt * 8))))
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color.lstrip("#").upper())
        borders.append(element)
    tbl_pr.append(borders)


def set_cell_margins(table, *, padding_pt: float) -> None:
    tbl_pr = table._tbl.tblPr
    existing = tbl_pr.find(qn("w:tblCellMar"))
    if existing is not None:
        tbl_pr.remove(existing)
    margins = OxmlElement("w:tblCellMar")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), _twips(padding_pt))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    tbl_pr.append(margins)


def shade_cell(cell, fill: str) -> None:
    if not fill:
        return
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill.lstrip("#").upper())
    tc_pr.append(shading)


def mark_header_row(row) -> None:
    """跨页重复：w:tblHeader。"""
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def _write_cell(cell, text: str, role_style) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    paragraph = cell.paragraphs[0]
    format_paragraph(paragraph, role_style)
    lines = str(text).split("\n")
    run = paragraph.add_run(lines[0])
    apply_role_font(run, role_style)
    for line in lines[1:]:
        run.add_break()
        extra = paragraph.add_run(line)
        apply_role_font(extra, role_style)


def _rectangular_rows(table: Table) -> list[list[str]]:
    """把行宽补齐/截断到列数 —— 上游多一列少一列都不允许渲染成歪表。"""
    width = table.column_count
    if width <= 0:
        return []
    return [(list(row) + [""] * width)[:width] for row in table.rows]


def add_table(doc, table: Table, style: StyleSpec) -> None:
    """把一个 Table 模型渲染到文档里（可选 caption + 真实表格）。"""
    if table.caption:
        add_paragraph(doc, table.caption, resolve(style, Role.TABLE_CAPTION))

    rows = _rectangular_rows(table)
    header = list(table.headers)[: table.column_count] if table.headers else []
    width = table.column_count
    if width <= 0 or (not header and not rows):
        return

    data_rows = rows
    total_rows = len(data_rows) + (1 if header else 0)
    doc_table = doc.add_table(rows=total_rows, cols=width)
    doc_table.alignment = _TABLE_ALIGN[style.table.align]
    doc_table.autofit = True

    table_style: TableStyle = style.table
    set_table_borders(
        doc_table, width_pt=table_style.border_width_pt, color=table_style.border_color
    )
    set_cell_margins(doc_table, padding_pt=table_style.cell_padding_pt)

    cursor = 0
    if header:
        header_row = doc_table.rows[0]
        for column, text in enumerate(header):
            _write_cell(header_row.cells[column], text, resolve(style, Role.TABLE_HEADER))
            shade_cell(header_row.cells[column], table_style.header_background)
        if table_style.repeat_header:
            mark_header_row(header_row)
        cursor = 1

    cell_role = resolve(style, Role.TABLE_CELL)
    for offset, row_values in enumerate(data_rows):
        row = doc_table.rows[cursor + offset]
        for column, text in enumerate(row_values):
            _write_cell(row.cells[column], text, cell_role)
