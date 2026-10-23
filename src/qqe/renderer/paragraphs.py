"""段落落盘工具：把 RoleStyle 翻译成 docx 段落格式。

行距统一用"倍数"语义（1.5 倍行距），与前端 CSS 的 line-height 一一对应，
这样预览与导出件的松紧感一致。
"""

from __future__ import annotations

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Pt

from qqe.domain.style import Align
from qqe.renderer.fonts import apply_role_font
from qqe.renderer.roles import RoleStyle

_ALIGN_MAP = {
    Align.LEFT: WD_ALIGN_PARAGRAPH.LEFT,
    Align.CENTER: WD_ALIGN_PARAGRAPH.CENTER,
    Align.RIGHT: WD_ALIGN_PARAGRAPH.RIGHT,
    Align.JUSTIFY: WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def format_paragraph(paragraph, role_style: RoleStyle) -> None:
    fmt = paragraph.paragraph_format
    fmt.alignment = _ALIGN_MAP[role_style.align]
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = role_style.line_spacing
    fmt.space_before = Pt(role_style.space_before_pt)
    fmt.space_after = Pt(role_style.space_after_pt)
    fmt.left_indent = Pt(role_style.left_indent_pt())
    fmt.first_line_indent = Pt(role_style.first_line_indent_pt())


def add_paragraph(container, text: str, role_style: RoleStyle):
    """新增一个段落（继承样式），文本按角色字体写入。"""
    paragraph = container.add_paragraph()
    format_paragraph(paragraph, role_style)
    if text:
        apply_role_font(paragraph.add_run(text), role_style)
    return paragraph
