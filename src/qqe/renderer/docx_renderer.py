"""DocxRenderer：QuestionSet + StyleSpec → .docx

三条铁律：
1. **无状态、无"增量"概念**：每次导出都从项目全量题目重新渲染。
   续加只发生在数据层；往 .docx 里追加会让新老样式混排。
2. **只认 `Question.order`**：渲染顺序、显示的题号都来自 order；
   `source_label`（原图题号）**绝不参与排序，也绝不出现在导出件里**。
3. **样式零硬编码**：所有字体/字号/间距/缩进都来自 StyleSpec，
   通过 roles.resolve() 落到段落角色上。
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt

from qqe.domain.models import Question, QuestionSet
from qqe.domain.style import PAPER_MM, StyleSpec
from qqe.renderer.fonts import apply_role_font
from qqe.renderer.paragraphs import add_paragraph
from qqe.renderer.roles import Role, resolve
from qqe.renderer.table_renderer import add_table
from qqe.renderer.text import chunk_options, option_line_text, question_text


class RenderError(ValueError):
    """数据不满足渲染前置条件。"""


def ordered_questions(question_set: QuestionSet) -> list[Question]:
    """按 order 升序返回；order 未分配/重复时**报错而不是猜**。

    顺序唯一真相是列表位置，order 是它的派生值。若拿到的 order 是 0 或重号，
    说明调用方漏了 renormalize() —— 这时候静默按列表顺序渲染，
    会让"用户拖动过的顺序"悄悄失效，宁可炸出来。
    """
    questions = list(question_set.questions)
    if not questions:
        return []
    for question in questions:
        if question.order <= 0:
            raise RenderError("存在未编号的题目（order=0），渲染前必须先 renormalize()")
    orders = [question.order for question in questions]
    if len(set(orders)) != len(orders):
        raise RenderError("题目 order 存在重复，渲染前必须先 renormalize()")
    return sorted(questions, key=lambda question: question.order)


def _setup_page(document, style: StyleSpec) -> None:
    section = document.sections[0]
    width_mm, height_mm = PAPER_MM[style.page.paper.value]
    section.page_width = Mm(width_mm)
    section.page_height = Mm(height_mm)
    section.top_margin = Cm(style.page.margin_top_cm)
    section.bottom_margin = Cm(style.page.margin_bottom_cm)
    section.left_margin = Cm(style.page.margin_left_cm)
    section.right_margin = Cm(style.page.margin_right_cm)


def _setup_default_font(document, style: StyleSpec) -> None:
    """把 Normal 样式也设成目标字体：避免空段/新段落回落到等线。"""
    normal = document.styles["Normal"]
    normal.font.name = style.font.latin
    normal.font.size = Pt(style.font.size_pt)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), style.font.latin)
    rfonts.set(qn("w:hAnsi"), style.font.latin)
    rfonts.set(qn("w:eastAsia"), style.font.cjk)
    rfonts.set(qn("w:cs"), style.font.latin)


def _write_header_footer(document, style: StyleSpec) -> None:
    section = document.sections[0]
    title_style = resolve(style, Role.TABLE_CELL)

    if style.page.header_text:
        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        apply_role_font(paragraph.add_run(style.page.header_text), title_style)

    if style.page.footer_page_number:
        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_page_number_field(paragraph, title_style)


def _add_page_number_field(paragraph, role_style) -> None:
    run = paragraph.add_run()
    apply_role_font(run, role_style)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


def _render_question(document, question: Question, style: StyleSpec) -> None:
    """一道题：题干 → 表格 → 选项。本期不渲染答案与解析。"""
    add_paragraph(document, question_text(question, style), resolve(style, Role.QUESTION))

    for table in question.tables:
        add_table(document, table, style)

    option_style = resolve(style, Role.OPTION)
    for chunk in chunk_options(question.options, style.option.options_per_line):
        add_paragraph(document, option_line_text(chunk, style), option_style)


class DocxRenderer:
    """渲染器是无状态的：同一样式、同一数据，永远得到同一份文件。"""

    def __init__(self, style: StyleSpec) -> None:
        self.style = style

    def render(self, question_set: QuestionSet) -> bytes:
        document = Document()
        _setup_page(document, self.style)
        _setup_default_font(document, self.style)
        _write_header_footer(document, self.style)

        if question_set.title:
            add_paragraph(document, question_set.title, resolve(self.style, Role.DOC_TITLE))

        current_group: str | None = None
        for question in ordered_questions(question_set):
            if question.group_title and question.group_title != current_group:
                add_paragraph(
                    document,
                    question.group_title,
                    resolve(self.style, Role.SECTION_TITLE),
                )
                current_group = question.group_title
            _render_question(document, question, self.style)

        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def render_to(self, question_set: QuestionSet, path: Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(self.render(question_set))
        return path


def render_docx(question_set: QuestionSet, style: StyleSpec) -> bytes:
    return DocxRenderer(style).render(question_set)
